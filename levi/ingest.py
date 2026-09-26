"""Turn an uploaded file into a Document, then into citeable Chunks.

Chunks keep character offsets into the document text, so we can
(a) map each chunk to a page for citations, and
(b) check retrieval against gold answer spans during eval.
"""
import hashlib
import io
import re
import unicodedata
from bisect import bisect_right
from dataclasses import dataclass, field
from pathlib import Path

from levi.schemas import Chunk

SUPPORTED_TYPES = {".pdf", ".docx", ".txt"}
_WORD = re.compile(r"\S+")


class UnsupportedFileType(ValueError):
    pass


class EmptyDocument(ValueError):
    pass


@dataclass
class Document:
    doc_id: str
    text: str
    # Char offset where each page starts. Empty for formats without pages.
    page_starts: list[int] = field(default_factory=list)

    def page_of(self, offset: int) -> int | None:
        if not self.page_starts:
            return None
        return bisect_right(self.page_starts, offset)  # 1-based page number


def normalize(text: str) -> str:
    # NFKC folds ligatures like "ﬁ" -> "fi" (these broke keyword search in v1's corpus).
    text = unicodedata.normalize("NFKC", text)
    return text.replace("\x00", "")


def content_hash(data: bytes) -> str:
    # Same bytes -> same doc_id, which later lets us skip re-embedding repeat uploads.
    return hashlib.sha256(data).hexdigest()[:16]


def load_bytes(data: bytes, filename: str) -> Document:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_TYPES:
        raise UnsupportedFileType(f"{ext or 'no extension'} not supported; use one of {sorted(SUPPORTED_TYPES)}")

    page_starts: list[int] = []
    if ext == ".pdf":
        from pypdf import PdfReader

        parts: list[str] = []
        offset = 0
        for page in PdfReader(io.BytesIO(data)).pages:
            page_text = normalize(page.extract_text() or "")
            page_starts.append(offset)
            parts.append(page_text)
            offset += len(page_text) + 2  # the "\n\n" joiner below
        text = "\n\n".join(parts)
    elif ext == ".docx":
        import docx

        text = normalize("\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs))
    else:
        text = normalize(data.decode("utf-8", errors="replace"))

    if not text.strip():
        # Scanned PDFs land here. OCR is on the roadmap, not in v2's scope.
        raise EmptyDocument("No extractable text (scanned PDFs need OCR, which is not supported yet)")
    return Document(doc_id=content_hash(data), text=text, page_starts=page_starts)


def from_text(text: str, doc_id: str | None = None) -> Document:
    text = normalize(text)
    return Document(doc_id=doc_id or content_hash(text.encode()), text=text)


def chunk_document(doc: Document, chunk_words: int, overlap: int) -> list[Chunk]:
    """Sliding window over words. Chunk text is an exact slice of doc.text."""
    if not 0 <= overlap < chunk_words:
        raise ValueError("overlap must be >= 0 and < chunk_words")

    spans = [(m.start(), m.end()) for m in _WORD.finditer(doc.text)]
    chunks: list[Chunk] = []
    step = chunk_words - overlap
    for start in range(0, len(spans), step):
        window = spans[start : start + chunk_words]
        char_start, char_end = window[0][0], window[-1][1]
        idx = len(chunks)
        chunks.append(
            Chunk(
                id=f"{doc.doc_id}:{idx}",
                doc_id=doc.doc_id,
                index=idx,
                text=doc.text[char_start:char_end],
                char_start=char_start,
                char_end=char_end,
                page=doc.page_of(char_start),
            )
        )
        if start + chunk_words >= len(spans):
            break
    return chunks
